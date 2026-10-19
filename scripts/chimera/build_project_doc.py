"""Builds the VIT course project document (Project Doc.pdf template, BCSE306L) as an editable .docx.

Run with the scratch venv:  data/scratch/docs_venv/bin/python scripts/chimera/build_project_doc.py
Then export to PDF:         soffice --headless --convert-to pdf --outdir docs/paper/project_doc docs/paper/project_doc/Project_Document.docx
Every result number is read from stored result files (data/eval/chimera/*.json); prose states the provenance. No human-annotated
result is used or claimed anywhere in this document.
"""
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path("data/eval/chimera")
OUTDIR = Path("docs/paper/project_doc")
FIG = OUTDIR / "figures"
J = lambda p: json.loads((ROOT / p).read_text())  # noqa: E731
pooled, dose, v5, adv, v6 = J("pd_final_pooled.json"), J("pd_dose_response.json"), J("pd_v5/analysis.json"), J("pd_adversarial_checks.json"), J("pd_v6/analysis.json")
plan = J("confirm_runpod_plan.json")
l2u = J("l2u_analysis.json")["difference_in_differences"]
pts = lambda x, d=1: f"{100 * x:+.{d}f}"  # noqa: E731
pc = lambda x, d=1: f"{100 * x:.{d}f}"  # noqa: E731
ci = lambda lo, hi: f"[{100 * lo:.1f}, {100 * hi:.1f}]"  # noqa: E731

BLUE, DRED = RGBColor(0x1F, 0x6F, 0xB5), RGBColor(0xA6, 0x1C, 0x1C)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
st = doc.styles["Normal"]
st.font.name, st.font.size = "Arial", Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
st.paragraph_format.space_after = Pt(5)
st.paragraph_format.line_spacing = 1.12


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def add_field(par, instr):
    r = par.add_run()
    for kind, text in (("begin", None), (None, instr), ("end", None)):
        if kind:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), kind)
        else:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = text
        r._r.append(e)
    return r


def rich(par, text, size=None, color=None, italic=False):
    """**bold** and *italic* inline markers."""
    for tok in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+`)", text):
        if not tok:
            continue
        if tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = par.add_run(tok[1:-1])
            r.font.name = "Courier New"
            r.font.size = Pt((size or 10) - 0.5)
            continue
        if tok.startswith("**"):
            r = par.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1])
            r.italic = True
        else:
            r = par.add_run(tok)
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
        if italic:
            r.italic = True
    return par


def P(text, align=None, size=None, after=None, italic=False, color=None):
    p = doc.add_paragraph()
    rich(p, text, size=size, italic=italic, color=color)
    if align == "j":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "c":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def H(num, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(10), Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"[{num}]  {text}")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(12.5), BLUE
    return p


def SH(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(6), Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold, r.font.size = True, Pt(10.5)
    return p


def BUL(items, size=None):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        rich(p, it, size=size)


def TABLE(header, rows, widths=None, size=8.5, head_fill="DCE6F1", caption=None):
    if caption:
        c = doc.add_paragraph()
        c.paragraph_format.keep_with_next = True
        c.paragraph_format.space_after = Pt(3)
        rich(c, caption, size=9)
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        rich(cell.paragraphs[0], f"**{h}**", size=size)
        shade(cell, head_fill)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            lines = str(v).split("\n")
            for k, ln in enumerate(lines):
                par = cells[i].paragraphs[0] if k == 0 else cells[i].add_paragraph()
                par.paragraph_format.space_after = Pt(1)
                rich(par, ln, size=size)
    if widths:
        t.autofit = False
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        t._tbl.tblPr.append(lay)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def FIGURE(name, caption, width=16.0):
    path = FIG / name
    if path.exists():
        doc.add_picture(str(path), width=Cm(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = P(caption, align="c", size=8.5, italic=True)
        p.paragraph_format.space_after = Pt(8)
    else:
        P(f"[MISSING FIGURE FILE {name}: {caption}]", color=DRED)


# ---------------------------------------------------------------- header and footer
hdr = sec.header.paragraphs[0]
hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = hdr.add_run("Project Document")
r.bold, r.font.size, r.font.name, r.font.color.rgb = True, Pt(20), "Cambria", DRED
ftr = sec.footer.paragraphs[0]
ftr.text = ""
ftr.add_run("AI Assignment\t\tC1 + TC1\t\tPage ").font.size = Pt(9)
add_field(ftr, "PAGE").font.size = Pt(9)

# ---------------------------------------------------------------- [1] title page
box = doc.add_table(rows=1, cols=1)
box.style = "Table Grid"
box.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = box.rows[0].cells[0]
cell.width = Cm(5)
cell.text = ""
rich(cell.paragraphs[0], "**[ Insert the VIT logo here ]**", size=11)
cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
for _ in range(2):
    doc.add_paragraph()
P("**Vellore Institute of Technology**", align="c", size=16)
P("**School of Computer Science and Engineering** [confirm the school name for your programme]", align="c", size=11)
P("Course Code and Name: **BCSE306L  Artificial Intelligence**     Slot: **C1 + TC1**", align="c", size=11)
for _ in range(2):
    doc.add_paragraph()
P("**Project Title**", align="c", size=11, after=2)
P("**Do Language Models Peek Ahead? Measuring Name-Keyed Memory Leakage When LLMs Read Serialised Fiction, on a Reader-Scoped Narrative Knowledge Graph (EchoTales)**", align="c", size=17)
for _ in range(2):
    doc.add_paragraph()
TABLE(["Member", "Registration Number"], [["Adhyan Jain", "24BCI0111"], ["Atharva Sheerh Pandey", "24BCI0069"], ["Aryan Gahlot", "24BCI0033"]], widths=[8, 6], size=11)
P("Submission: 16 October 2026   |   Repository: EchoTales (branch chimera/system-adapters)", align="c", size=9)
doc.add_page_break()

# ---------------------------------------------------------------- [2] abstract
H(2, "Structured Abstract")
lu_all, vl_all = pooled["luna_only"]["all_v2_v3_v6"], pooled["valid_pool"]["all_v2_v3_v6"]
rec = v5["pools"]["G2_valid"]
d_lu = dose["luna_gap_vs_local_recognition"]
abstract = (
    "**Background:** Language models recap and answer questions about serialised fiction for readers who have read only part of a book. A model that remembers the book may name a disguised character's true identity before the book reveals it. "
    "**Objective:** To measure that leakage, test whether name memory rather than lost name cues causes it, and audit a narrative knowledge graph for the same property. "
    "**Methods:** From public-domain novels we built 38 events in which a character known by an alias is later named; the book's own reveal sentence is the oracle. Each mention was shown with original names, with a consistent letter cipher, and with cue-preserving pseudonyms. The leakage gap is the drop in premature naming under renaming, with event-level statistics. Five local 7-9B models and one commercial reference model were tested; a truncate-and-rebuild audit tested the graph pipeline. "
    f"**Results:** The reference model named the identity early more often with original names: gap {pts(lu_all['gap_A_minus_B'])} points (95% CI {ci(*lu_all['ci'])}; {lu_all['events_positive']} events positive, {lu_all['events_negative']} negative). Four local models pooled {pts(vl_all['gap_A_minus_B'])} ({ci(*vl_all['ci'])}) but did not replicate on seven fresh books. "
    f"Pseudonyms restored control accuracy, not premature naming (recovered fraction {rec['recovered_fraction']:+.2f}). The gap rose with book familiarity (Spearman {d_lu['spearman']:+.2f}, p = {d_lu['perm_p']:.3f}). The graph resolver was prefix-causal (0 of 28,165 comparisons); its persona stage looked ahead for 8.0% of characters. "
    "**Conclusions:** Name-keyed memory leakage is measurable and not explained by lost name cues for the strongest model. Events are model-selected, no human labels are used, and a pre-registered confirmatory study is pending."
)
nwords = len(re.sub(r"\*\*", "", abstract).split())
assert nwords <= 250, nwords
P(abstract, align="j")
P(f"*({nwords} words)*", size=8, after=2)

# ---------------------------------------------------------------- [3] introduction
H(3, "Introduction")
P("**Background and problem.** Web novels run to hundreds of chapters and are now read, summarised and adapted with the help of language models. Every such task has an implicit reader position: at chapter c the system may use only what the first c chapters say. Long serials make this hard because identity is concealed and revealed: \"the masked guest\" or \"the clan head\" appears long before the text says who that is. Standard tooling gives one omniscient answer per character [11, 10]. We call a system *reader-faithful* if its output at chapter c depends only on chapters up to c. Two things can break that. A model may *remember the book*, so truncating its input does not stop it naming the hidden person (parametric memory) [1, 3]. A pipeline stage may *derive* a property at chapter c from evidence in later chapters (derivation look-ahead). Figure 6 shows both.", align="j")
P("**Objective and methodology.** The project measures the first violation and audits the second. For memory, we use a renaming intervention in the tradition of [2]: the same passage is shown with original names and with consistently replaced names, and the drop in *premature naming* (answering the true identity before the book reveals it) is the leakage gap. Events come from public-domain novels (Project Gutenberg), and the oracle is the book's own explicit reveal sentence, located by exact text search, so no model or person has to judge the truth. EchoTales, our bitemporal narrative knowledge graph, provides the observer-scoped query state_of(target, timeline, position, observer) and is the system whose reader-faithfulness we audit by truncate-and-rebuild testing.", align="j")
P("**Method discipline.** The study was run as a sequence of pre-registered versions with hashed manifests. Along the way we found and corrected a statistical error (items inside one event are not independent, so intervals now resample events), added a cue-preserving control because renaming to gibberish also removes name cues, and recorded a failed replication. A larger confirmatory study with open-weight models on rented GPUs is fully specified and hashed but has not been run; nothing in this document depends on it.", align="j")
P(f"**Results and analysis.** For the reference model the leakage gap is {pts(lu_all['gap_A_minus_B'])} points over 38 events, and it appears in all three datasets. Four local 7-9B models show {pts(vl_all['gap_A_minus_B'])} pooled but zero on the seven least-known books. Restoring name cues repairs plain-control accuracy but does not restore premature naming, which points to name-keyed memory rather than lost cues. The gap rises with how well the book is recognised. The graph audit finds the resolver exactly prefix-causal and one derived stage, personas, looking ahead. Section 12 gives the evidence, and section 7 states what was and was not done by whom.", align="j")

# ---------------------------------------------------------------- [4] related work
H(4, "Background / Related Work")
P("The team studied the works below. Every research article is from 2020 or later; the tools and web resources are listed separately. Provenance of our reading is recorded in the repository (docs/research/NOVELTY_REPORT/PRIOR_ART.md and PRIOR_ART_E0.md): articles marked in the table were read in their own text, not from abstracts.", align="j")
P("**Memorisation of books by language models.** Chang et al. [1] probe which books ChatGPT and GPT-4 know with a name-cloze task over 571 works and show that memorisation tracks how often passages appear on the web and that memorised books score better on downstream tasks. Fu et al. [5] compare summaries of 25 well-known books written from the model's memory with summaries written from the full text and find the full text is usually, but not always, better. Both measure recall or summary quality; neither asks what a reader at a given chapter is entitled to know, and neither controls names in a reading task.", align="j")
P("**Renaming and anonymisation as controls.** Jiang and Ferraro [2] replace character names (masking, same-culture and gender-preserving cross-culture replacement) and report accuracy drops of up to 45.2% on character-understanding benchmarks, with human validation that gold labels survive. SynthWorlds [7] contrasts real and synthetic parallel worlds to separate knowledge from reasoning, with a rename-leak audit. Deußer et al. [4] study pseudonymisation with consistent placeholders and mapping the answer back, but as a privacy-utility trade-off over eleven benchmarks, not as a memorisation control. None of them uses a reveal time, and none separates memory from lost name cues.", align="j")
P("**Time-aware and disclosure-aware narrative state.** TimeChara [3] defines point-in-time character hallucination in role-play and lists spoiler avoidance as a motivation, but defines leakage as contradiction with a timeline. Ma et al. [6] name premature information disclosure as one failure of interactive narrators on movie synopses without testing memory. Narrative World Model [14] and NarraWorld [15] model chapter-safe memory and reveal order, and E2RAG [16] builds temporally consistent entity-event graphs; none of them measures memory-driven leakage of an identity.", align="j")
P("**Coreference and character resolution in fiction.** LitBank [11], KoCoNovel [10], BOOKCOREF [9] and LlmLink [8] supply corpora and systems for literary coreference. KoCoNovel is notable for annotating both reader and omniscient views, but its models are tested on the omniscient version. Web-novel resources such as GenWebNovel [12] contain only the first ten chapters of each novel, so later reveals are structurally absent, and Song and Liu [13] list hidden identity as an unsolved challenge in a single sentence.", align="j")
P("**Tools and resources.** BookNLP, vLLM, ollama, Project Gutenberg, RunPod, Hugging Face and the Claude Code command line supply the software, data and compute used here (Table 2).", align="j")
rw = [
    ["1", "Chang et al. [1]\nEMNLP, ACL, 2023", "• Name-cloze probe of memorisation\n• Memorisation tracks web duplication (Spearman 0.74, pre-1928 books)\n• Memorised books do better on downstream tasks", "571 works (name-cloze passages)", "• Tests recall, not a reading task with a reveal\n• Closed models only; no renamed control on downstream use"],
    ["2", "Jiang and Ferraro [2]\narXiv:2412.14368 v6, 2026", "• Name replacement as a memorisation constraint\n• Masked, same-culture, gender-preserving cross-culture names\n• Accuracy drop up to 45.2%", "Six character-understanding benchmarks", "• No chapter prefix, reveal time or alias task\n• Does not split the drop into memory versus lost cues"],
    ["3", "Ahn et al. [3]\nFindings of ACL, 2024", "• Point-in-time role-play hallucination benchmark\n• About 10,895 instances\n• Narrative-Experts decomposition as a mitigation", "Characters from four book series", "• Leakage = timeline contradiction, not tied to memory\n• No renaming control; four series only"],
    ["4", "Deußer et al. [4]\narXiv:2609.11335, 2026", "• Pseudonymisation with mapping back, plus four other anonymisers\n• 11 benchmarks, five models (7B-72B)\n• Pseudonymisation least harmful to utility (as summarised in our notes)", "Eleven public benchmarks", "• Privacy-utility framing; memorisation untested\n• No narrative reveal task"],
    ["5", "Fu et al. [5]\narXiv:2603.09981, 2026", "• Summaries from memory versus full text\n• LLM judges\n• Full text usually better; memory sometimes wins", "25 well-known public-domain books", "• No reader-position cutoff\n• LLM judges carry self-preference risk"],
    ["6", "Ma et al. [6]\narXiv:2608.08160, 2026", "• Long-horizon consistency benchmark with a fact ledger\n• Defines premature information disclosure\n• Six narrator models", "100 CMU Movie Summary synopses", "• No memory or renaming test (read via summary tool)\n• Famous films confound memory with reasoning"],
    ["7", "Gu et al. [7]\nICLR, 2026", "• Parallel real and synthetic worlds\n• Knowledge-advantage gap\n• Rename-leak audit", "Wikidata-derived parallel corpora (multi-hop QA)", "• Encyclopaedic QA, not fiction\n• No alias or reveal task"],
    ["8", "Zhu et al. [8]\nCOLING, 2025", "• Dual-LLM linking with memorisation\n• Chunked long narratives\n• Evaluated with MUC, B3, CoNLL", "LitBank and DOC-ARC", "• Static identity, no time axis\n• Stops at the graph; no leakage control"],
    ["9", "Martinelli et al. [9]\nACL, 2025", "• Book-scale coreference corpus\n• Book-scale silver and gold sets\n• BookNLP CoNLL-F1 42.2 on gold", "Gutenberg books (PG19-derived)", "• Omniscient labels only\n• Web availability acknowledged, not controlled"],
    ["10", "Kim et al. [10]\nLREC-COLING, 2024", "• Reader and omniscient coreference views\n• Three annotators, MUC agreement 94.53\n• Views differ in 5 of 50 works", "KoCoNovel: 50 Korean novels, 178K tokens", "• Experiments on omniscient view only\n• Korean; no LLM or leakage control"],
    ["11", "Bamman et al. [11]\nLREC, 2020", "• Coreference annotation of literature\n• Revelation of identity discussed (Sec. 4.2.5)", "LitBank: 100 English works", "• Annotates textual reality, not the reader's state of knowledge\n• Omniscient labels only"],
    ["12", "Zhao et al. [12]\nCOLING, 2025", "• Nested NER for Chinese web novels\n• PER, LOC, ORG\n• Coreference left as future work", "400 chapters (first 10 of 40 novels), 1.2M tokens", "• Later reveals structurally absent\n• Chinese; data-use agreement restricts reuse"],
    ["13", "Song and Liu [13]\nLREC-COLING, 2024", "• Resolving character representations in a Chinese novel\n• Zero and non-zero anaphora\n• Names hidden identity as open problem", "One Chinese print novel", "• One qualitative sentence on hidden identity\n• No metric, LLM or prefix setup"],
    ["14", "Saifullah et al. [14]\narXiv:2607.05577, 2026", "• Writer memory with chapter-safe cutoffs\n• Reveal order versus event order\n• Typed temporal state", "Long-form fiction memory QA", "• No identity task or leakage control\n• Scores memory QA accuracy only"],
    ["15", "Tian et al. [15]\narXiv:2609.32401, 2026", "• Validity-interval story memory\n• Alias-linking stage\n• Premature disclosure stated as a design aim", "ConStory-Bench and continuation checkpoints", "• Premature disclosure never measured\n• Alias linking not evaluated"],
    ["16", "Zhang et al. [16]\narXiv:2506.05939, 2025", "• Entity-event knowledge graph for RAG\n• Temporal-causal consistency\n• 'Never merge mentions' design", "ChronoQA: 497 QA pairs over 9 Gutenberg stories", "• Avoids merging rather than resolving\n• No leakage control"],
]
TABLE(["Sl#", "Articles", "Work Done", "Dataset", "Gaps / Limitations"], rw, widths=[0.9, 3.0, 5.4, 3.6, 4.3], size=7.5, caption="**Table 1.** Research articles reviewed (2020 onward).")
tools = [
    ["T1", "BookNLP (Bamman, GitHub)", "Character clustering, coreference and speaker attribution for books; baseline family", "Omniscient output; no reader-position or leakage notion found"],
    ["T2", "vLLM (Kwon et al., SOSP 2023)", "PagedAttention serving engine; OpenAI-compatible server used on the rented GPU", "Throughput depends on GPU and quantisation; to be measured in calibration"],
    ["T3", "ollama (ollama.com)", "Local model runtime used for the 7-9B models on the 8 GB laptop GPU", "Small models only; run-to-run determinism not guaranteed"],
    ["T4", "Project Gutenberg (gutenberg.org)", "Public-domain texts, the data source for all events", "Pre-1928 Western fiction; texts may be in training data"],
    ["T5", "RunPod pricing page (dated 2026-10-08)", "GPU hourly prices used for the confirmatory cost plan", "Prices change; read through a summary tool, so approximate"],
    ["T6", "Hugging Face model cards (OLMo-2-32B, Llama-3.3-70B AWQ, Qwen2.5-72B AWQ)", "Licences, quantisation and serving notes for the confirmatory models", "Community 4-bit builds may differ from full precision"],
    ["T7", "Claude Code CLI v2.1.301, headless mode", "Programmatic access to Claude Sonnet 5.5 through a subscription (measured overhead 814 tokens, 1.9 s per call)", "No temperature or seed control; usage limits; terms for bulk use unverified"],
    ["T8", "Wikidata API", "Planned objective fame variable (number of Wikipedia editions) for the dose-response", "Coverage follows the Western canon"],
]
TABLE(["Sl#", "Tool / report", "Purpose in this project", "Gaps / Limitations"], tools, widths=[0.9, 4.4, 7.0, 4.9], size=7.5, caption="**Table 2.** Blogs, reports and tools studied.")

# ---------------------------------------------------------------- [5] gaps
H(5, "Summary of Limitations / Gaps")
gaps = [
    ["G1", "Memorisation studies measure recall (name cloze, source identification, summaries) but not premature disclosure at a reader position [1, 2, 5].", "C1"],
    ["G2", "Renaming controls turn names into gibberish or other cultures; none separates name-keyed memory from lost name cues, and none checks that the control itself is valid per model [2, 4, 7].", "C2, C3"],
    ["G3", "Time-aware and disclosure-aware narrative systems model chapter-safe state but do not measure memory-driven leakage of an identity or the effect of a repair [3, 6, 14, 15, 16].", "C5, C6"],
    ["G4", "Literary coreference resources use omniscient labels or early chapters, so reveals are absent or invisible [9, 10, 11, 12, 13].", "C1"],
    ["G5", "Statistical practice in the sections we read treats items as independent and reports single accuracies; replication and pre-registration are rare (our own first analysis made the same error and corrected it).", "C4"],
    ["G6", "Pseudonymisation is studied for privacy and utility, not as a mitigation of leakage with a measured utility cost on post-reveal answers [4].", "C6"],
]
TABLE(["Gap", "Limitation found in the literature", "Addressed by"], gaps, widths=[1.2, 13.0, 2.6], size=8.5, caption="**Table 3.** Gaps identified from the literature review.")

# ---------------------------------------------------------------- [6] contributions
H(6, "Primary Contributions of the Project")
BUL([
    "**C1. A reveal-time leakage measure with a text-derived oracle (gaps G1, G4).** Premature naming is scored against the book's own reveal sentence found by exact search, with a verifier (cutoff, no lexical cue, controls) and an independent model audit of events. 38 events from 38 public-domain books were analysed; a 150-event set is specified.",
    f"**C2. A cue-preserving pseudonym control and the recovered-fraction decomposition (gap G2).** Gender- and name-type-matched pseudonyms restore plain-control accuracy (e.g. qwen2.5 80.3 to 96.7) but not premature naming: gap {pts(rec['gap_A_minus_B3'])} against {pts(rec['gap_A_minus_B'])}, recovered fraction {rec['recovered_fraction']:+.2f} {ci(*rec['recovered_fraction_ci'])}.",
    "**C3. A memory-free validity gate (gap G2).** Comparing two independent renamings with each other, plus renaming on unseen synthetic text, checks the control without original names; it showed that plain controls on famous books are themselves assisted by memory or by name cues (e.g. gemma2: 23 control items right only with original names, 1 only renamed; sign test p = 3e-6).",
    f"**C4. Event-level inference, honest replication and a familiarity dose-response (gap G5).** Item-level intervals were replaced by event-cluster intervals; the failed replication on seven fresh books is reported; the gap rises with book recognition (Spearman {d_lu['spearman']:+.2f}, p = {d_lu['perm_p']:.3f}, cross-family).",
    "**C5. A prefix-causality audit of a narrative graph pipeline (gap G3).** Truncate-and-rebuild testing found the resolver exactly causal (0 of 28,165 mention comparisons) and the persona stage looking ahead for 8.0% of character checks (23 of 288), and framed both leakage channels as violations of reader-faithfulness.",
    "**C6. A pre-registered confirmatory design with a mitigation and a harm task (gaps G3, G6).** 150 events, four open and closed model families, pseudonymise-and-map-back as the mitigation with a 5-point utility margin, and a recap-spoiler task; power, cost, hashes and stopping rules are fixed. It has not been run.",
    "**C7. The substrate: EchoTales, a bitemporal observer-scoped narrative graph** with the query state_of(target, timeline, position, observer), a resolver and generation consumers, whose defects the audit located and fixed (Track S, engineering).",
])

# ---------------------------------------------------------------- [7] individual
H(7, "Individual Contributions")
ind = [
    ["Adhyan Jain\n24BCI0111", "**Lead for the research and the system.**\n• Defined the research question, the leakage measure and the study design; wrote every pre-registration (PD v1 to v7 and the confirmatory study) with hashed manifests.\n• Architected EchoTales: bitemporal observer-scoped graph, state_of, resolver, persona and world-context stages (system architecture, full pipeline implementation); rendering and voice integration.\n• Wrote the experiment and analysis code (about 50 scripts): event builder, letter-cipher and cue-preserving pseudonym renamers, local and API model runners, validity gate, event-cluster bootstrap, dose-response, label audit; 1,320 automated tests pass.\n• Ran all experiments (local 7-9B models on the laptop GPU; one commercial reference model with a spend ledger of about USD 0.43) and diagnosed the failures (validity gate versus engine; statistical error in item-level intervals).\n• Built the evaluation harness, baselines and ablations; ran the prefix-causality audit and the engineering fixes.\n• Designed the confirmatory study, its power and cost model and the GPU deployment plan; produced this document, the slides and the demo."],
    ["Atharva Sheerh Pandey\n24BCI0069", "• Annotation and labelling responsibilities, test design and verification (roles as recorded in the project repository).\n• Shared the literature study required by the course.\n• Human annotation: assigned role; **no human-annotated labels exist yet and none are used in any result here**."],
    ["Aryan Gahlot\n24BCI0033", "• Annotation and labelling responsibilities, test design and verification (roles as recorded in the project repository).\n• Shared the literature study required by the course.\n• Human annotation: assigned role; **no human-annotated labels exist yet and none are used in any result here**."],
]
TABLE(["Member", "Contribution (implementation and documentation)"], ind, widths=[3.4, 13.4], size=8.5, caption="**Table 4.** Individual contributions. *(Team to confirm the wording of the second and third rows before submission.)*")
P("Human labelling of the 72-item pilot sheet and of the confirmatory events is planned as shared work and will be reported only in the journal version, once it has been done.", size=9, italic=True)

# ---------------------------------------------------------------- [8] tools
H(8, "Tools and Technologies Used")
TABLE(["Category", "Item", "Use"], [
    ["Hardware", "Laptop, NVIDIA RTX 4060 8 GB GPU; about 15 GiB RAM", "All local model runs, analysis and the demo"],
    ["Hardware (planned)", "One RunPod pod, H100 or A100 80 GB, 250 GB disk", "Confirmatory study on 70B-class open models (cost plan in section 12)"],
    ["Language and data", "Python 3.12 (uv), SQLite with an explicit bitemporal schema", "EchoTales graph, resolver, evaluation harness"],
    ["Models run so far", "qwen2.5:7b, llama3.1:8b, mistral:7b-instruct, gemma2:9b, phi4-mini (ollama); gpt-5.6-luna (API reference model)", "Exploratory and pre-registered experiments (PD v2 to v7)"],
    ["Models planned", "Llama-3.3-70B and Qwen2.5-72B (4-bit AWQ), OLMo-2-32B (public training data), Claude Sonnet 5.5 (CLI)", "Confirmatory study; vLLM on RunPod"],
    ["Analysis", "NumPy-free standard-library statistics (bootstrap, permutation, sign test), matplotlib, python-docx, python-pptx", "Results, figures, documents"],
    ["Quality", "pytest (1,320 tests), ruff, pre-registrations hashed with SHA-256", "Reproducibility"],
    ["Interface", "React review interface and a stdlib HTTP API", "Research view for the demonstration (section 14)"],
], widths=[3.2, 8.0, 5.6], size=8.5, caption="**Table 5.** Hardware and software.")

# ---------------------------------------------------------------- [9] methodology
H(9, "Methodology Used")
SH("9.1 Events and the text-derived oracle")
P("An *event* is a character referred to by one or more alias expressions (\"the Opera ghost\") who is later named in an explicit sentence of the book (\"Erik\"). The reveal sentence is located by whitespace-insensitive exact search; the cutoff is a chapter heading, or a fixed position, before it; the verifier requires at least three alias mentions before the cutoff, no sentence before it containing both alias and identity, no shared name token, and a cast list of 7 to 9 names. An independent model audit labels each reveal EXPLICIT, IMPLIED or NOT SUPPORTED. The oracle fixes what the *book states*, not what a reader might infer from hints [10, 11].", align="j")
SH("9.2 Items, prompts and arms")
P("For each event we sample up to eight mentions of the alias before the cutoff, mentions after the reveal, and plain controls (ordinary cast names). Each item shows the last 1,200 characters ending at the mention's sentence, the marked expression, and the cast list in a seeded order with UNDETERMINED as an option; the model answers with one letter as JSON. *Premature naming* is answering the true identity before the reveal. Arms: **A** original names; **B** all proper nouns replaced by a letter cipher; **B2** a second independent cipher (validity checks); **B3** natural pseudonyms matched in gender and name type for the event's cast, other proper nouns ciphered, so place names stay memory-free. Window text and candidates are identical across arms.", align="j")
SH("9.3 Measures")
P("The **leakage gap** is the premature-naming rate in A minus the mean in B (and B2). The **recovered fraction** (rate B3 minus rate B) over (rate A minus rate B) is the share of the gap that restoring name cues brings back. The **validity gate** G2 requires agreement of B with B2 on plain controls of at least 90% and of A with B on unseen synthetic items of at least 80%. **Memory probes** are a closed-book alias question and a title-assisted question (title and alias, no passage).", align="j")
SH("9.4 Statistics")
P("Because items within an event share one book and one reveal, the **event** is the resampling unit: event-cluster bootstrap intervals [17], exact sign tests on events, permutation tests for correlations [18], Holm correction across models [19], and DerSimonian-Laird random-effects pooling with Hartung-Knapp intervals for the confirmatory study [20, 21]. Wilson intervals are used for proportions [22].", align="j")
SH("9.5 Pre-registration and replication")
P("Rules, models, prompts and decision thresholds were written and hashed before the data they govern were collected (PD v1 to v7; a confirmatory study is specified in the same way) [23]. Several verdicts were INCONCLUSIVE or NOT SUPPORTED and are reported as such.", align="j")
SH("9.6 Graph audit (reader-faithfulness of the pipeline)")
P("The pipeline is run on chapters 1..k only and on the full volume; the facts visible to the reader observer at chapter k are compared (truncate-and-rebuild). For the resolver this was done at 14 cuts over three serials; for the persona stage at 288 character-cuts. The serials are copyrighted and stay local; only counts are reported.", align="j")
SH("9.7 Confirmatory design (specified, not run)")
P("150 events (30 development, 120 test, split by book) from at least 60 new books; four model families; the mitigation (pseudonymise names, run, map names in the answer back) with a 5-percentage-point utility margin on post-reveal and control accuracy and a plain 'use only the passage' prompt as baseline; a recap-and-profile spoiler task; 40 LLM-written unseen stories as the unseen-text control. Power: 80% power to detect a gap of 9.0 points (Holm-adjusted) with 120 test events.", align="j")

# ---------------------------------------------------------------- [10] dataset
H(10, "Dataset Used")
P("**Source.** Project Gutenberg [24], public-domain English-language fiction published before 1928 (some translated from French, Spanish or Russian). Texts are downloaded locally and are not stored in the repository; items store offsets and hashes. Event selection was made by language-model drafters and verified by exact search; **no event has been verified by a person**, and **no human annotation is used** anywhere in this document. Events are model-selected, so every number is labelled *text-derived*.", align="j")
TABLE(["Dataset", "Books / events", "Items", "Used for"], [
    ["PD v2", "19 books, 19 events", "317 (135 pre-reveal, 30 post-reveal, 152 controls)", "Repaired design; 6 models"],
    ["PD v3", "12 fresh books, 12 events", "216 (96 / 24 / 96)", "Confirmatory test, 80% validity bar"],
    ["PD v6", "7 fresh books, 7 events", "124 (54 / 14 / 56)", "Final pre-registered test (not replicated for local models)"],
    ["PD v7", "9 books (5 famous, 4 obscure), 9 events", "157 (72 / 13 / 72)", "Prospective fame contrast; local models only; inconclusive"],
    ["L2U / G2b", "Synthetic stories with invented names", "see l2u_analysis.json", "Unseen-text control and validity gate"],
    ["Local serials", "3 web novels (RI 199, LOTM 40, ORV 40 chapters)", "11,174 mentions (RI)", "Graph audits only; copyrighted, never distributed"],
], widths=[2.4, 4.6, 5.4, 4.4], size=8.5, caption="**Table 6.** Datasets (all event sets overlap no other set).")
P("**Sample event.** Book: *The Phantom of the Opera* (Project Gutenberg 175). Alias expressions: \"the Opera ghost\", \"the ghost\". Revealed identity: *Erik*. Cast list shown to the model: seven to nine names including Christine, Raoul de Chagny, the Persian and Erik. A pre-reveal item shows the 1,200 characters before one mention of the alias, with the expression marked; the correct reader answer is UNDETERMINED, while answering Erik is premature. Characteristics: pre-reveal items are sampled at distances from a few hundred to tens of thousands of characters before the reveal; controls are mentions of ordinary cast members, used to check that renaming did not damage reading.", align="j", size=9.5)

# ---------------------------------------------------------------- [11] approach
H(11, "Proposed Approach")
FIGURE("fig_arch.png", "**Figure 5.** System architecture of the leakage study (ten stages, from corpus to statistics).", width=16.4)
P("**Process, stage by stage.** (1) Public-domain novels are the corpus. (2) Event mining finds sentences with reveal cues and extracts alias and identity. (3) A deterministic verifier and an independent audit keep only explicit, clean events. (4) The event list, cast lists and the dev/test split by book are frozen and hashed before any model is run. (5) The item builder cuts the 1,200-character prefix window and fixes candidate lists. (6) Each item is shown in four arms with identical windows and candidates; the memory probes run per event. (7) Models answer: local models through ollama, open models through vLLM on a RunPod pod, and one closed model through the Claude CLI. (8) Scoring marks an answer premature if it is the true identity before the reveal. (9) Statistics are fixed in advance: event-cluster bootstrap, random-effects pooling, permutation dose-response, Holm correction. (10) Outputs are the gap, the recovered fraction, the mitigation and spoiler tests, and the report.", align="j")
FIGURE("fig_channels.png", "**Figure 6.** Reader-faithfulness and the two leakage channels measured in this project.", width=14.5)

# ---------------------------------------------------------------- [12] results
H(12, "Result Analysis and Discussion")
P("All results below are *text-derived* (the book's own reveal sentence) unless marked; none is human-annotated. Intervals are 95% event-cluster intervals. The local runs used 7-9B models; the reference model is a commercial API model (total spend about USD 0.43, ledgered).", size=9, italic=True)
FIGURE("fig_forest.png", "**Figure 1.** Leakage gap (points) by model group and dataset. The dotted line marks +10 points. The seven-book replication (v6) is positive for the reference model and not for the pooled local models.", width=15.5)
mrows = []
short = {"gpt-5.6-luna": "gpt-5.6-luna (API)", "gemma2:9b": "gemma2 9B", "llama3.1:8b": "llama3.1 8B", "mistral:7b-instruct": "mistral 7B", "phi4-mini:latest": "phi4-mini", "qwen2.5:7b": "qwen2.5 7B"}
for m in ("gpt-5.6-luna", "gemma2:9b", "llama3.1:8b", "mistral:7b-instruct", "phi4-mini:latest", "qwen2.5:7b"):
    c2, c3, c6 = adv["pd_v2"]["per_model_cluster"][m], adv["pd_v3"]["per_model_cluster"][m], pooled["v6_per_model"][m]["cluster"]
    mrows.append([short[m], f"{pts(c2['pooled_gap'])} {ci(*c2['event_cluster_ci95'])}", f"{pts(c3['pooled_gap'])} {ci(*c3['event_cluster_ci95'])}", f"{pts(c6['gap_A_minus_B'])} {ci(*c6['ci'])}"])
TABLE(["Model", "v2: 19 events", "v3: 12 fresh events", "v6: 7 fresh events"], mrows, widths=[3.8, 4.4, 4.4, 4.2], size=8.5, caption="**Table 7.** Leakage gap per model and dataset (points, 95% event-cluster interval).")
P(f"**Reading Figure 1 and Table 7.** The reference model's gap is positive in all three datasets ({pts(lu_all['gap_A_minus_B'])} points over all 38 events, {lu_all['events_positive']} events positive against {lu_all['events_negative']} negative, sign test p = {lu_all['sign_p']:.0e}). The four valid local models together give {pts(vl_all['gap_A_minus_B'])} {ci(*vl_all['ci'])} over 38 events (sign test p = {vl_all['sign_p']:.2f}), but on the seven least-known books (v6) their pooled gap is {pts(pooled['valid_pool']['v6']['gap_A_minus_B'])} {ci(*pooled['valid_pool']['v6']['ci'])}: the pre-registered replication **failed** for the small models and held for the reference model. Section 12 therefore supports a narrower claim than 'models leak': the strongest model leaks clearly, and small models leak where they know the book.", align="j")
FIGURE("fig_pooled.png", "**Figure 4.** Pooled over all 38 events with event-level inference.", width=11.5)
FIGURE("fig_cue.png", "**Figure 2.** (a) Plain-control accuracy under original names (A), the letter cipher (B) and cue-preserving pseudonyms (B3); (b) leakage gap against the cipher and against pseudonyms.", width=16.0)
crow = []
for m in ("gemma2:9b", "qwen2.5:7b", "mistral:7b-instruct", "llama3.1:8b", "phi4-mini:latest", "gpt-5.6-luna"):
    x = v5["per_model"][m]["pd_v2"]
    crow.append([short[m], f"{pc(x['ctrl_acc_A'])} / {pc(x['ctrl_acc_B'])} / {pc(x['ctrl_acc_B3'])}", f"{pts(x['per_event']['gap_A_minus_B'])}", f"{pts(x['per_event']['gap_A_minus_B3'])} {ci(*x['per_event']['gap_A_minus_B3_ci'])}"])
TABLE(["Model", "Control accuracy % (A / B / B3)", "Gap A minus B", "Gap A minus B3 [95%]"], crow, widths=[3.6, 5.0, 3.0, 5.2], size=8.5, caption="**Table 8.** Cue-preserving control on the 19 v2 events (controls and gaps).")
P(f"**Reading Figure 2 and Table 8.** Replacing names with a letter cipher lowers control accuracy by 5 to 15 points, so name-borne cues matter for reading. Natural pseudonyms bring accuracy back to the original level, yet the premature-naming gap does not shrink: pooled over 31 events it is {pts(rec['gap_A_minus_B'])} {ci(*rec['gap_A_minus_B_ci'])} with the cipher and {pts(rec['gap_A_minus_B3'])} {ci(*rec['gap_A_minus_B3_ci'])} with pseudonyms (recovered fraction {rec['recovered_fraction']:+.2f}, interval {ci(*rec['recovered_fraction_ci'])}; the pre-registered 'memory survives' line was 1/3). What original names carry that pseudonyms do not is the specific name, which is what memory is keyed on.", align="j")
FIGURE("fig_dose.png", "**Figure 3.** Mean leakage gap by how many local models recognise the book (38 events; post hoc).", width=12.8)
P(f"**Reading Figure 3.** For the four valid local models the gap is about zero on books that at most two of the five recognise and large (+40.9 points) on books all five recognise (Spearman {dose['valid_local_pool_gap_vs_local_recognition']['spearman']:+.2f}, permutation p = {dose['valid_local_pool_gap_vs_local_recognition']['perm_p']:.5f}); this explains the failed replication, because the seven new books are poorly recognised. The cross-family check, the reference model's gap against the local models' recognition, is weaker but positive (Spearman {d_lu['spearman']:+.2f}, p = {d_lu['perm_p']:.3f}). Caveats: post hoc; the local-model correlation shares models with its predictor; cell sizes are small.", align="j")
SH("Graph audit and error corrections")
TABLE(["Check", "Result", "Provenance"], [
    ["Resolver prefix-causality (truncate-and-rebuild, 14 cuts, 3 serials)", "0 of 28,165 mention-cut pairs differ (ids stable); the by-design whole-volume passes do break it (up to 39% of early mentions in one serial)", "System-internal, no labels (eng_probe_E1_*.json and ENGINEERING_NOVELTY_PROBE/SUMMARY.md)"],
    ["Persona stage prefix-causality (288 character-cuts)", "23 differ (8.0%); gender is set from later pronoun evidence and body splits come from a whole-volume rule", "System-internal, no labels (ENGINEERING_NOVELTY_PROBE/SUMMARY.md)"],
    ["Resolver versus plain string clustering", "Adjusted Rand index 0.971 to 0.985 against surface-string clustering", "System-internal; precision not measured"],
    ["Statistical error found", "Item-level bootstrap treated items of one event as independent; intervals widened after switching to event resampling (e.g. v3 alone: lower bound 1.3 points)", "Corrected in the analysis"],
], widths=[5.6, 6.4, 4.8], size=8, caption="**Table 9.** Graph audit and corrections.")
SH("Run times (measured)")
TABLE(["Run", "Calls", "Local 7-9B models (minutes per model)", "Reference API model"], [
    ["PD v2, three arms plus probes", "1,132", "9 to 16 (qwen 11.5, llama 10.9, mistral 9.0, phi4 9.5, gemma2 16.1)", "about 20 min; USD 0.12"],
    ["PD v3", "808", "6 to 11 (llama 7.2, mistral 5.8, phi4 6.0, gemma2 10.6)", "USD 0.085"],
    ["PD v5, cue-preserving arm", "533", "4 to 9 (mistral 4.3, phi4 4.9, qwen 6.5, llama 6.6, gemma2 9.0)", "9.3 min; USD 0.059"],
    ["PD v6", "538", "4 to 9 (qwen 6.0, llama 6.0, mistral 4.2, phi4 5.6, gemma2 8.7)", "9.4 min; USD 0.057"],
    ["PD v7 (local only)", "525", "4 to 7 (qwen 4.9, llama 5.0, mistral 4.1, phi4 4.6, gemma2 7.0)", "cancelled"],
    ["Live probe in the demo (three arms, one item)", "3", "a few seconds on an 8 GB GPU (measured in the demo, section 14)", "-"],
], widths=[4.2, 1.5, 7.6, 3.5], size=8, caption="**Table 10.** Wall-clock times from the run logs.")
SH("Confirmatory study: planned cost and time")
cp = plan["core_150_events_7_items"]
h100 = cp["3_open_models"]["H100_80GB_community_assuming_2.5x"]
a100 = cp["3_open_models"]["A100_80GB_community"]
P(f"The pre-registered confirmatory study (150 events, 7 items each; Llama-3.3-70B, Qwen2.5-72B, OLMo-2-32B, plus Claude Sonnet 5.5 through the CLI) needs about {h100['total_hours']} GPU-hours on an H100 (USD {h100['expected_usd']}, worst case {h100['worst_2x_usd']}) or about {a100['total_hours']} hours on an A100 (USD {a100['expected_usd']}); the H100 figure assumes a 2.5x speed-up that a 25-minute calibration run will measure first. These are plans, not results.", align="j", size=9.5)
SH("Discussion")
P("The evidence supports a measured, qualified claim: name-keyed memory leakage is clear for the strongest model tested, is not explained by lost name cues, rises with book familiarity, and is weak or absent for small models on books they do not know. It does not show that EchoTales or any web-serial resolver leaks, and it does not use human labels. The main weaknesses are stated openly: events are model-selected, the memory probes give the book title, the pseudonyms are heuristic, the pooled interval for local models sits at the edge of significance, and the replication failed for small models.", align="j")

# ---------------------------------------------------------------- [13] comparative
doc.add_page_break()
H(13, "Comparative Analysis")
TABLE(["Work", "Task and control", "Reveal time?", "Memory vs lost cues?", "Their headline number", "Ours (same direction)"], [
    ["Chang et al. [1]", "Name cloze on 571 works (recall)", "No", "No", "Memorisation correlates with web hits (Spearman 0.74, pre-1928 books)", f"Familiarity correlates with leakage gap: {d_lu['spearman']:+.2f} (cross-family), {dose['valid_local_pool_gap_vs_local_recognition']['spearman']:+.2f} (local pool)"],
    ["Jiang and Ferraro [2]", "Six character tasks; renamed characters", "No", "No (gender-preserving variant, no decomposition)", "Up to 45.2% accuracy drop", f"Premature-naming gap {pts(lu_all['gap_A_minus_B'])} points (reference model)"],
    ["Ahn et al. [3] TimeChara", "Role-play at a story time", "Yes (timeline)", "No", "State-of-the-art models hallucinate future facts", "Leakage tied to memory by renaming"],
    ["Deußer et al. [4]", "Pseudonymise, map back; 11 benchmarks", "No", "No", "Pseudonymisation least harmful", "Mitigation evaluated for leakage and utility (design only)"],
    ["Gu et al. [7] SynthWorlds", "Real versus synthetic worlds", "No", "No", "Knowledge-advantage gap", f"Real-minus-synthetic difference-in-differences {pts(l2u['D'])} {ci(*l2u['ci95'])} (does not exclude 0)"],
    ["KoCoNovel [10]", "Reader and omniscient labels", "Annotated, not tested on models", "No", "Views differ in 5 of 50 works", "Text-derived reveal oracle instead of annotators"],
], widths=[2.5, 3.2, 1.7, 2.3, 3.4, 3.7], size=7.5, caption="**Table 11.** Comparison with related work. Numbers are not directly comparable: tasks, corpora and models differ.")
P("**What the comparison shows.** The direction of our findings matches the closest prior work (renaming lowers performance on famous books; memorisation follows how widely a book circulates), which supports the validity of the measurement. What is added is (i) a reveal-time task with a text-derived oracle, (ii) a test that separates memory from lost name cues, (iii) a validity check that does not rely on original names, (iv) event-level inference and a reported failed replication, and (v) a reader-faithfulness audit of a graph pipeline. We did not re-run any of these systems; the comparison is of design and reported numbers, not of a shared benchmark.", align="j", size=9.5)

# ---------------------------------------------------------------- [14] user manual
H(14, "User Manual")
P("**Command line (reproducing the reported numbers from stored results; no GPU needed).**", size=9.5)
BUL([
    "`source .venv/bin/activate` then `python scripts/chimera/pd_final_pooled.py`, `python scripts/chimera/pd_dose_response.py`, `python scripts/chimera/pd_v5_cuepres.py analyze`: recompute Tables 7 and 8 and Figures 1 to 4 from the stored run files in data/eval/chimera/.",
    "`data/scratch/docs_venv/bin/python scripts/chimera/make_project_figures.py` regenerates the figures; `build_project_doc.py` regenerates this document.",
    "`python -m pytest packages/core/tests packages/pipeline/tests` runs the 1,320 tests.",
], size=9)
P("**Research view of the review interface (demonstration).** Start the interface as described in docs/paper/project_doc/FRONTEND_RESEARCH_VIEW.md and open the Research page. It has (1) headline cards, (2) the gap by model and dataset, (3) the cue-preserving panel, (4) the familiarity dose-response and (5) a live probe that runs one public-domain item under original names, the letter cipher and pseudonyms on a local model and shows the three answers with timings.", align="j", size=9.5)
for i, (fn, cap) in enumerate([("ui_01_research_overview.png", "**Figure 7.** Research page: overview cards."), ("ui_02_gap_by_model.png", "**Figure 8.** Gap by model and dataset."), ("ui_03_cue_and_dose.png", "**Figure 9.** Cue-preserving panel and dose-response."), ("ui_04_live_probe_result.png", "**Figure 10.** Live probe result with per-arm timings.")]):
    FIGURE(fn, cap, width=15.0)

# ---------------------------------------------------------------- [15] references
H(15, "References")
refs = [
    "Chang, Kent K., Mackenzie Cramer, Sandeep Soni, and David Bamman. \"Speak, Memory: An Archaeology of Books Known to ChatGPT/GPT-4.\" *Proceedings of the 2023 Conference on Empirical Methods in Natural Language Processing*, Association for Computational Linguistics, 2023. arXiv:2305.00118.",
    "Jiang, Yuxuan, and Francis Ferraro. \"Beyond Math: Stories as a Testbed for Memorization-Constrained Reasoning in LLMs.\" arXiv:2412.14368 v6, 24 Apr. 2026.",
    "Ahn, Jaewoo, et al. \"TimeChara: Evaluating Point-in-Time Character Hallucination of Role-Playing Large Language Models.\" *Findings of the Association for Computational Linguistics: ACL 2024*, 2024. arXiv:2405.18027.",
    "Deußer, et al. \"On the Impact of Anonymization on the Performance of LLMs.\" arXiv:2609.11335 v2, Sep. 2026. [first names to be completed from the arXiv page]",
    "Fu, et al. \"Reading or Remembering\" (book summarisation from memory versus full text). arXiv:2603.09981, Feb. 2026. [full title and first names to be completed from the arXiv page]",
    "Ma, et al. Long-horizon consistency benchmark for interactive narratives (premature information disclosure). arXiv:2608.08160, Aug. 2026. [full title and first names to be completed from the arXiv page]",
    "Gu, Ken, et al. \"SynthWorlds: Controlled Parallel Worlds for Disentangling Reasoning and Knowledge in Language Models.\" *International Conference on Learning Representations (ICLR)*, 2026. arXiv:2510.24427.",
    "Zhu, et al. \"LlmLink: Dual LLMs for Dynamic Entity Linking on Long Narratives with Collaborative Memorisation and Prompt Optimisation.\" *Proceedings of COLING 2025*, 2025.",
    "Martinelli, Giuliano, et al. \"BOOKCOREF: Coreference Resolution at Book Scale.\" *Proceedings of the 63rd Annual Meeting of the Association for Computational Linguistics*, 2025. arXiv:2507.12075.",
    "Kim, Kyuhee, Surin Lee, and Sangah Lee. \"KoCoNovel: Annotated Dataset of Character Coreference in Korean Novels.\" *Proceedings of LREC-COLING 2024*, 2024. arXiv:2404.01140.",
    "Bamman, David, Olivia Lewke, and Anya Mansoor. \"An Annotated Dataset of Coreference in English Literature.\" *Proceedings of the 12th Language Resources and Evaluation Conference (LREC)*, 2020.",
    "Zhao, et al. \"GenWebNovel\" (nested named-entity corpus of Chinese web novels). *Proceedings of the 31st International Conference on Computational Linguistics (COLING 2025)*, 2025, pp. 3836-3849. [full title to be completed]",
    "Song, et al. Resolving different representations of fictional characters in a Chinese novel. *Proceedings of LREC-COLING 2024*, 2024, pp. 1408-1421. [full title to be completed]",
    "Saifullah, Mohammad, et al. \"Narrative World Model.\" arXiv:2607.05577, Jul. 2026. [subtitle to be completed]",
    "Tian, Qiuyu, et al. \"Shared Worlds, Private Minds.\" arXiv:2609.32401, Sep. 2026. [subtitle to be completed]",
    "Zhang, Ze Yu, et al. E2RAG and ChronoQA: entity-event knowledge graphs for retrieval-augmented generation. arXiv:2506.05939, 2025. [full title to be completed]",
    "Efron, Bradley, and Robert J. Tibshirani. *An Introduction to the Bootstrap.* Chapman and Hall, 1993.",
    "Dror, Rotem, et al. \"The Hitchhiker's Guide to Testing Statistical Significance in Natural Language Processing.\" *Proceedings of the 56th Annual Meeting of the Association for Computational Linguistics*, 2018.",
    "Holm, Sture. \"A Simple Sequentially Rejective Multiple Test Procedure.\" *Scandinavian Journal of Statistics*, vol. 6, no. 2, 1979, pp. 65-70.",
    "DerSimonian, Rebecca, and Nan Laird. \"Meta-Analysis in Clinical Trials.\" *Controlled Clinical Trials*, vol. 7, no. 3, 1986, pp. 177-188.",
    "Hartung, Joachim, and Guido Knapp. \"On Tests of the Overall Treatment Effect in Meta-Analysis with Normally Distributed Responses.\" *Statistics in Medicine*, vol. 20, no. 12, 2001, pp. 1771-1782.",
    "Wilson, Edwin B. \"Probable Inference, the Law of Succession, and Statistical Inference.\" *Journal of the American Statistical Association*, vol. 22, no. 158, 1927, pp. 209-212.",
    "Nosek, Brian A., et al. \"The Preregistration Revolution.\" *Proceedings of the National Academy of Sciences*, vol. 115, no. 11, 2018, pp. 2600-2606.",
    "Project Gutenberg. *Project Gutenberg.* gutenberg.org. Accessed 16 Oct. 2026.",
    "Kwon, Woosuk, et al. \"Efficient Memory Management for Large Language Model Serving with PagedAttention.\" *Proceedings of the 29th ACM Symposium on Operating Systems Principles (SOSP)*, 2023.",
    "RunPod. \"Pricing.\" runpod.io/pricing (page dated 8 Oct. 2026). Accessed 16 Oct. 2026.",
    "Hugging Face. Model repositories casperhansen/llama-3.3-70b-instruct-awq, Qwen/Qwen2.5-72B-Instruct-AWQ, allenai/OLMo-2-0325-32B-Instruct. huggingface.co. Accessed 16 Oct. 2026.",
    "Anthropic. Claude Code command-line interface, version 2.1.301 (`claude --help`, print mode). Accessed 16 Oct. 2026.",
]
for i, r in enumerate(refs, 1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(0.8), Cm(-0.8)
    p.paragraph_format.space_after = Pt(2)
    rich(p, f"[{i}]  {r}", size=8.5)
P("*Items [1] to [16] are the reviewed articles (2020 onward); [17] to [23] support the statistical and pre-registration methods; [24] to [28] are tools and web resources. Details marked in square brackets could not be confirmed from the project's reading notes and must be completed from each paper's page before submission.*", size=8)

# ---------------------------------------------------------------- [16] appendix
H(16, "Appendix")
P("**a) Link to the presentation (PPT):** [paste the shareable link after uploading EchoTales_Research_Presentation.pptx]", size=9.5)
P("**b) Link to the pre-recorded demonstration video:** [paste the link after recording; script in Demo_Video_Script.docx]", size=9.5)
P("**c) Source files and steps to execute:** [paste the repository link after you push the branch; the commits are local until then]. Steps are in docs/paper/project_doc/RUN_STEPS.md and summarised here:", size=9.5)
BUL([
    "Clone, then `uv sync` (Python 3.12); `source .venv/bin/activate`.",
    "Re-download the public-domain texts into data/raw/pd_gutenberg/ (Project Gutenberg ids listed in docs/research/preregistrations/*events*.json).",
    "Recompute results from stored runs: `python scripts/chimera/pd_final_pooled.py` and `python scripts/chimera/pd_dose_response.py` (seconds, no GPU).",
    "Re-run a model: `python scripts/chimera/pd_v6.py run --texts data/raw/pd_gutenberg_v6 --model qwen2.5:7b` (about 6 minutes on an 8 GB GPU with ollama running; this overwrites that model's stored result file, so copy data/eval/chimera/pd_v6 first).",
    "Start the interface and open the Research page (FRONTEND_RESEARCH_VIEW.md).",
], size=9)
P("**Provenance statement.** All numbers are text-derived, synthetic or system-internal. No human-annotated result is reported. Event selection is model-drafted and not yet verified by a person; the confirmatory study and human labelling are future work.", size=9, italic=True)

OUTDIR.mkdir(parents=True, exist_ok=True)
doc.save(OUTDIR / "Project_Document.docx")
print("saved", OUTDIR / "Project_Document.docx", "abstract words", nwords)
